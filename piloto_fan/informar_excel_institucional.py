"""Produce informe y diccionario agregados, locales, sin modificar la aplicación."""
from pathlib import Path
import json
import html
import hashlib
import pandas as pd
from explorar_excel_institucional import OUT, ROOT, SOURCE


def read(name):
    return json.loads((OUT / name).read_text(encoding="utf-8"))


def fmt(x):
    return f"{x:,}".replace(",", ".") if isinstance(x, int) else str(x)


def decimal(x, digits=3):
    return f"{x:.{digits}f}".replace(".", ",")


def table(headers, rows):
    return '<div class="scroll"><table><thead><tr>' + "".join(
        f"<th>{html.escape(str(c))}</th>" for c in headers) + "</tr></thead><tbody>" + "".join(
        "<tr>" + "".join(f"<td>{html.escape(fmt(c))}</td>" for c in row) + "</tr>"
        for row in rows) + "</tbody></table></div>"


p, d, s = read("perfil_completo.json"), read("diagnostico_uniones_y_cobertura.json"), read("alcance_fan_propuesto.json")
fan, core = s["scopes"]["FAN_etiquetada"], s["scopes"]["FAN_propuesta_presencial"]
previous_dir = ROOT / "piloto_fan/artifacts"
last = json.loads((previous_dir / "latest.json").read_text())["run_id"]
previous = json.loads((previous_dir / last / "report.json").read_text(encoding="utf-8"))
assert hashlib.file_digest(SOURCE.open("rb"), "sha256").hexdigest() == p["sha256"]
assert p["quality"]["grades_all"]["numeric"] == 550860
assert p["quality"]["grades_fan"]["numeric"] == 64018
assert sum(fan["coded_grades"].values()) + fan["numeric_grade_rows"] == fan["history_rows"]
assert p["sheets"]["h"]["columns"] + p["sheets"]["m"]["columns"] == 52

descriptions = {
    "PIDM": "Identificador interno para vincular las dos hojas. No es MRUN; excluir como predictor y de exportaciones públicas.",
    "TERM_CODE": "Código del período observado. Usar calendario por régimen; el orden numérico no acredita disponibilidad.",
    "CAMP_CODE": "Código del campus. Diferenciar campus del plan y campus de la asignatura.",
    "LEVL_CODE": "Código del nivel: PR corresponde a Pregrado y PO a Pregrado Online en este extracto.",
    "NRC": "Identificador de sección/curso. No usarlo solo ni asumir catálogo constante; existen inconsistencias por período.",
    "SUBJ_CODE": "Código de materia o área de la asignatura; conservar junto al número de curso.",
    "CRSE_NUMB": "Número/código de asignatura. No interpretar como cantidad ni como dificultad.",
    "SECCION": "Sección informada; 356 filas sin valor.",
    "CREDIT_HOURS": "Magnitud de carga declarada: unidad pendiente de confirmar. No tratar como SCT ni ponderar todavía.",
    "GRDE_CODE_FINAL": "Resultado final como texto: nota con coma decimal o código institucional. Conservar ambas representaciones.",
    "MODO_CALIF": "Modo de calificación; E en todas las filas, sin variación predictiva.",
    "REPETIDO": "Indicador nominal de repetición, 100% vacío. No significa que no existan repitencias.",
    "CRSE_TITLE": "Nombre de la asignatura; etiqueta descriptiva, no clave única.",
    "PERIODO": "Nombre del período: semestres, trimestres, verano, TAV, bloques y causales.",
    "NIVEL": "Descripción de nivel. Pregrado y Pregrado Online deben poder analizarse separadamente.",
    "CAMPUS": "Descripción de sede/campus del plan; incluye Campus Virtual.",
    "STVAFDC_CODE": "Código asociado a jornada/modalidad en esta extracción.",
    "JORNADA": "Diurno, Vespertino, E-Learning o B-Learning. No sustituye el régimen de calendario.",
    "STST_CODE": "Código del estado del alumno. No equivale al estado del plan.",
    "STATUS_ALUMNO": "Estado general del alumno; casi todo Activo. No usarlo como etiqueta automática de permanencia.",
    "STSP_CODE": "Código del estado del plan; requiere diccionario y fecha de vigencia.",
    "STATUS_PLAN_ESTUDIO": "Activo, retirado, retracto, suspensión y otros. Es estado en el extracto, sin fecha del cambio.",
    "PROGRAM": "Código de programa/plan. Parte de la clave única de matrícula junto a PIDM y período.",
    "PROGRAMA": "Descripción del programa/plan. No confundir con código de carrera.",
    "MAJR_CODE": "Código de carrera. Preferible a uniones por nombres abreviados.",
    "CARRERA": "Nombre institucional de carrera; requiere catálogo y tratamiento de excepciones FAN.",
    "ACYR_CODE": "Año académico declarado; contrastar con calendario y período.",
    "RESD_CODE": "Código de categoría de residencia; no es comuna ni condición socioeconómica.",
    "RESD_DESC": "Residente / Extranjero. Uso descriptivo sujeto a pertinencia, no predictor por defecto.",
    "COLL_CODE": "Código de facultad/unidad académica; 62 valores ausentes en el universo.",
    "FACULTAD": "Facultad/unidad asociada al plan. Mejora el alcance FAN, pero contiene excepciones.",
    "DEGC_CODE": "Código del tipo de salida del programa; no acredita un título obtenido.",
    "TIPO_SALIDA": "Título Profesional, Licenciatura, Técnico o Sin Título ni Grado. No es fecha/resultado de titulación.",
    "TERM_CODE_CTLG": "Período de catálogo curricular. No asumir que es el ingreso ni el corte de extracción.",
    "TERM_CODE_MATRIC": "Período de matrícula declarado; 902 ausentes. Definición exacta pendiente.",
    "ADMT_CODE": "Código de vía de admisión, con 30,24% ausente en el universo.",
    "TIPO_ADMISION": "Vía de admisión; no contiene puntajes PAES, NEM ni ranking.",
    "CURR_RULE": "Identificador/regla curricular administrativa; no introducir como variable numérica continua.",
    "RATE_CODE": "Código de duración/modalidad administrativa de matrícula; validar semántica.",
    "DURACION_MATRICULA": "MATRICULA ANUAL o SEMESTRAL. No es duración nominal de la carrera.",
    "KEY_SEQNO": "Secuencia administrativa. No define por sí sola el plan prioritario o vigente.",
    "PRIORIDAD": "Prioridad administrativa; requiere regla de selección institucional. No elegir mínimo/máximo arbitrariamente.",
    "STYP_CODE": "Código del tipo de alumno, diferente del estado del alumno.",
    "TIPO_ALUMNO": "Nuevo, Continuo y otras categorías administrativas. No equivale automáticamente a primera vez en la UA.",
    "TERM_CODE_ADMIT": "Período de admisión registrado; no acredita que exista historia académica desde ese año.",
    "CMJR_RULE_1_1": "Regla de carrera/plan, identificador técnico con 13 ausentes.",
    "SECUENCIA_PLAN": "Secuencia del plan, útil para auditoría; no equivale sin confirmación a una fecha.",
    "SEQNO": "Secuencia administrativa; conservar para trazabilidad y resolución de múltiples planes.",
}
inventory = []
for label, sheet in p["sheets"].items():
    for col, profile in sheet["profile"].items():
        short = col.split("_", 1)[1]
        assert short in descriptions, short
        inventory.append({
            "hoja": "STG UWVHIAC" if label == "h" else "STG UWVMAHI",
            "variable": col, "tipo_leido": profile["dtype"],
            "ausentes": profile["missing"], "porcentaje_ausentes": profile["missing_pct"],
            "valores_distintos": profile["distinct"], "interpretacion": descriptions[short],
            "distribucion_o_rango": json.dumps(
                profile.get("values", profile.get("numeric", {})), ensure_ascii=False)})

# Describe annual membership without calling unobserved students dropouts.
m = pd.read_parquet(OUT / "m_local.parquet")
fm = m.loc[m.UWVMAHI_FACULTAD.str.contains("Negocios", case=False, na=False)]
first = set(fm.loc[fm.UWVMAHI_TERM_CODE.astype(str).str.startswith("2024"), "UWVMAHI_PIDM"])
next_fan = set(fm.loc[fm.UWVMAHI_TERM_CODE.astype(str).str.startswith("2025"), "UWVMAHI_PIDM"])
next_ua = set(m.loc[m.UWVMAHI_TERM_CODE.astype(str).str.startswith("2025"), "UWVMAHI_PIDM"])
continuity = {"base_fan_2024": len(first), "observed_fan_2025": len(first & next_fan),
              "observed_other_faculty_only_2025": len((first & next_ua) - next_fan),
              "not_observed_ua_2025": len(first - next_ua),
              "warning": "Presencia en el extracto, todos los estados. No son tasas oficiales de retención ni deserción."}
(OUT / "presencia_anual.json").write_text(json.dumps(continuity, ensure_ascii=False, indent=2), encoding="utf-8")

body = """
<header><p class="eyebrow">ANALÍTICA DE ESTUDIANTES · UNIVERSIDAD AUTÓNOMA DE CHILE</p>
<h1>Nueva fuente institucional: diagnóstico e incorporación al piloto FAN</h1>
<p>20 de octubre de 2026 · Uso interno · Exploración integral del Excel, sin entrenar nuevos modelos ni publicar datos.</p></header>
<div class="callout"><strong>Conclusión:</strong> el proyecto ya dispone de un extracto académico institucional real.
Puede avanzar hacia un piloto de desempeño y progresión entre períodos en Administración y Negocios.
Todavía no dispone de información semanal para alertas tempranas ni de una conexión operativa a Banner/Canvas.</div>
<nav><a href="#datos">Datos</a> · <a href="#fan">Alcance FAN</a> · <a href="#calidad">Calidad</a> ·
<a href="#aporte">Aporte</a> · <a href="#integracion">Integración</a> · <a href="#estado">Estado</a> ·
<a href="#plan">Plan</a> · <a href="#variables">52 campos</a></nav>
<h2 id="datos">1. Qué se revisó</h2>
<p>El usuario confirmó que el archivo fue generado por la universidad. La fecha de corte, el área emisora,
la consulta de extracción y los filtros de estados siguen pendientes de confirmación.
Los nombres técnicos son compatibles con un extracto académico institucional; no prueban por sí solos
un acceso directo a Banner. Se debe tratar como información interna, no como una nueva base pública.</p>
"""
body += table(["Hoja", "Filas de datos", "Columnas", "Personas distintas", "Períodos distintos"], [
    ["STG UWVHIAC · historia", p["sheets"]["h"]["rows"], 13, p["sheets"]["h"]["people"], len(p["sheets"]["h"]["periods"])],
    ["STG UWVMAHI · matrícula/planes", p["sheets"]["m"]["rows"], 39, p["sheets"]["m"]["people"], len(p["sheets"]["m"]["periods"])],
    ["Hoja1", 0, 0, 0, 0]])
body += f"""<p>Se recorrieron las 703.123 filas y los 52 campos de las hojas con datos. Se comprobaron faltantes,
cardinalidad, valores, duplicados, claves, uniones, calendarios y alcance FAN; también fórmulas, errores y enlaces
externos del libro: no se encontraron estos últimos tres elementos.
La fuente conserva su SHA-256: <code>{p['sha256']}</code>.</p>
<p>El alcance temporal observado es 2024–2025. No hay historia de notas de 2023 ni de 2026.
Las cohortes de admisión más antiguas no amplían esa historia. Matrícula distingue
Pregrado (131.371 filas) y Pregrado Online (2.604). El código PO significa Pregrado Online en este archivo.</p>
<h2 id="fan">2. Alcance de Administración y Negocios</h2>"""
body += table(["Universo", "Personas", "Filas de matrícula/plan", "Registros de asignaturas vinculados", "Personas con historia"], [
    ["Etiqueta FAN del extracto", fan["people"], fan["rows"], fan["history_rows"], fan["people_with_history"]],
    ["Propuesta de seis denominaciones en cuatro campus físicos", core["people"], core["rows"], core["history_rows"], core["people_with_history"]]])
body += """<p>La vinculación anterior es por persona y período: no acredita que cada asignatura pertenezca al plan
FAN cuando existen planes simultáneos. La segunda fila es un alcance propuesto, no una población aprobada
ni restringida exclusivamente a estados activos. Excluye Online, intercambio y la anomalía de Arquitectura;
mantiene las denominaciones originales sin inventar equivalencias entre carreras.</p>"""
body += table(["Carrera según el extracto FAN", "Personas distintas", "Filas"], [
    [x["UWVMAHI_CARRERA"], x["people"], x["rows"]] for x in fan["by_career"]])
body += table(["Campus · alcance propuesto", "Personas distintas"], [
    [x["UWVMAHI_CAMPUS"], x["people"]] for x in core["by_campus"]])
body += """<p>Las personas pueden figurar en más de una carrera, campus o nivel en la ventana;
por eso no corresponde sumar sus subtotales. La presencia de una fila de Arquitectura en FAN
debe revisarse; los intercambios y la modalidad online requieren decisiones explícitas de alcance.
Las dos denominaciones de Control de Gestión (D/T) permanecen separadas.</p>
<h2 id="calidad">3. Calidad: hallazgos que cambian el análisis</h2>"""
quality_rows = [
    ["Claves", "Sin PIDM ni períodos nulos. Sin filas idénticas duplicadas en ambas tablas.",
     "Historia es única por PIDM + período + NRC; matrícula por PIDM + período + programa."],
    ["Planes múltiples", "170 pares persona–período tienen 341 filas; 171 filas adicionales a ese grano. 86 pares abarcan facultades diferentes.",
     "Conservar plan y resolver atribución. No elegir arbitrariamente la primera fila ni la mayor secuencia."],
    ["Unión de tablas", "686 registros de historia no tienen matrícula en el mismo período; 11.777 persona–período de matrícula no tienen historia.",
     "Preservar y explicar no coincidencias. No convertir ausencia de historia en reprobación, inactividad o cero cursos inscritos."],
    ["Multiplicación silenciosa", "568.462 registros académicos tienen coincidencia; una unión directa produce 569.054 filas.",
     "El incremento de 592 filas quedaría casi oculto al perder las 686 sin coincidencia. Exigir validación de cardinalidad."],
    ["FAN y atribución", "36 pares persona–período FAN están en más de una facultad. En el alcance propuesto, 63.264 registros académicos tienen un único registro institucional de plan.",
     "Mantener el resto separado hasta confirmar la regla de atribución. Revisar también 255 diferencias entre campus del plan y del curso en ese subconjunto."],
    ["Notas", "550.860 notas numéricas entre 1 y 7; 18.288 resultados codificados. Coma decimal en el texto.",
     "Conservar nota original, nota numérica y código. No convertir RI/RR/RM/RC/CR/P/CC/REPROB automáticamente en una nota o un mismo tipo de reprobación."],
    ["Resultados FAN", "64.018 notas numéricas y 2.828 códigos: RI 2.719; RM 62; RR 46; CR 1.",
     "Los códigos representan 4,23% de sus registros vinculados. Una tasa que los ignore debe explicar su denominador."],
    ["Carga", "CREDIT_HOURS toma valores de 0 a 1.280; predominan 72, 108 y 36; también aparecen 3–10.",
     "Confirmar unidades y cambios por plan/régimen. No usar como créditos SCT ni ponderar promedios hasta resolverlo."],
    ["Repitencia", "REPETIDO está vacío en el 100% de las 569.148 filas.",
     "Solo pueden derivarse intentos repetidos observados dentro de 2024–2025, con catálogo de equivalencias; no repitencia histórica completa."],
    ["Admisión", "40.514 filas (30,24%) sin vía de admisión. Sin PAES, NEM, ranking ni información de financiamiento.",
     "No confundir descripción de admisión con puntaje. No unir PIDM a MRUN por similitud numérica ni de perfiles."],
    ["Fechas y estados", "Sin fechas de publicación de notas ni cambios de estado; 133.972 filas tienen alumno Activo, pero existen planes retirados, retractados y otros.",
     "Definir estados efectivos con VRA. Los estados actuales pueden contener información posterior al corte de una predicción retrospectiva."],
    ["Calendarios", "18 períodos de matrícula y 14 de historia: semestres, trimestres, verano, TAV, bloques y causales.",
     "Crear un calendario explícito por régimen. No comparar todos como semestres ni inferir fechas del orden numérico."],
    ["Catálogo de cursos", "Hay grupos período–NRC con más de un código, nombre o carga; añadir campus/nivel no elimina todos los conflictos.",
     "Conservar el detalle original, auditar conflictos y construir una dimensión de asignatura/sección con clave validada."],
    ["Otros faltantes", "Historia: 5 campus y 356 secciones. Matrícula: 902 períodos de matrícula, 95 categorías de residencia y 62 códigos de facultad.",
     "Mantener ausentes explícitos y mostrar cobertura por subgrupo; no imputar códigos administrativos sin evidencia."],
    ["Interpretaciones falsas", "TIPO_SALIDA es el tipo de título/grado del programa; DURACION_MATRICULA es anual/semestral.",
     "No prueban titulación efectiva ni duración nominal de carrera. Cuatro filas FAN tienen código de admisión mayor al período observado: revisar con calendario."],
]
body += table(["Dimensión", "Evidencia", "Tratamiento"], quality_rows)
body += """<h2>4. Dos hallazgos descriptivos ya utilizables</h2>
<p><strong>Desempeño:</strong> en el alcance FAN propuesto, excluyendo persona–período con múltiples registros
institucionales de plan y considerando solo semestres ordinarios:</p>"""
body += table(["Período", "Personas con historia", "Notas numéricas", "Notas bajo 4", "% entre notas numéricas"], [
    [x["UWVHIAC_TERM_CODE"], x["people"], x["numeric_grades"], x["below4"],
     f"{100*x['below4']/x['numeric_grades']:.2f}%".replace(".", ",")]
    for x in core["semester_grade_summary_unambiguous"]])
body += """<p>La proporción pasa de aproximadamente 7,0% a 9,8% entre el primer semestre de 2024 y el segundo de 2025.
Es una señal para investigar asignaturas, cohortes y cobertura; no demuestra deterioro causal, no compara la misma
composición de estudiantes y no es la tasa oficial de reprobación: los códigos no numéricos no están en el denominador.</p>"""
body += table(["Seguimiento de las 3.549 personas etiquetadas FAN en algún período de 2024", "Personas"], [
    ["Observadas en FAN en 2025", continuity["observed_fan_2025"]],
    ["Observadas en UA en 2025, solo fuera de FAN", continuity["observed_other_faculty_only_2025"]],
    ["No observadas en la matrícula UA del extracto 2025", continuity["not_observed_ua_2025"]]])
body += """<p>La última fila no equivale a deserción: puede incluir egreso/titulación, suspensión, problemas de cobertura
u otras salidas. El universo incluye todos los estados y años de ingreso; no es una cohorte de primer año.</p>
<h2 id="aporte">5. Qué aporta y qué todavía no permite</h2>"""
body += table(["Uso", "Situación con la nueva fuente"], [
    ["Tablero académico FAN", "Viable después de resolver reglas: resultados por asignatura, carrera, campus, jornada, régimen y período, con coberturas y denominadores."],
    ["Segmentación de desempeño", "Nueva posibilidad real: perfiles de rendimiento, heterogeneidad de notas, carga observada, intentos repetidos observados y progresión. Es distinta de la segmentación de ingreso existente."],
    ["Caso predictivo 1", "Propuesta: probabilidad de un resultado académico adverso en el próximo período ordinario, usando solo resultados de períodos cerrados. La etiqueta de reprobación requiere diccionario oficial."],
    ["Caso predictivo 2", "Propuesta: no renovación de matrícula en el siguiente período comparable. Requiere estados efectivos, población elegible y tratamiento de egreso/suspensión/cambios."],
    ["Student 360", "Viable una vista interna de trayectoria 2024–2025 seudonimizada, con planes, cursos y resultados; faltan trazas semanales y mecanismos institucionales de acceso/contacto."],
    ["Benchmark externo", "Se conserva con fuentes públicas para comparar universidades. No hay notas de otras universidades en el Excel, ni puente PIDM–MRUN."],
    ["Alertas semanales de reprobación/desenganche", "Todavía no: faltan notas parciales fechadas, asistencia, actividad Canvas y entregas."],
    ["Titulación", "El modelo público previo permanece como línea complementaria. Este Excel no contiene una tabla de títulos obtenidos ni sus fechas."],
    ["Operación en 2026", "No sustentada por un extracto limitado a 2024–2025. Solicitar actualización antes de emitir prioridades sobre alumnos actuales."],
])
body += """<p>Los dos nuevos casos deben priorizarse con las direcciones. La postulación pide segmentación,
dos casos priorizados y al menos tres alternativas por caso; no impone que el segundo sea desenganche Canvas.
El extracto permite acercarse al propósito original sin simular notas ni matrículas para entrenar o evaluar.</p>
<h2 id="integracion">6. Incorporación sin alterar la solución existente</h2>
<ol>
<li><strong>Recepción trazable:</strong> mantener el Excel original inmutable y registrar emisor, fecha de corte,
filtros, alcance, hash y responsable de actualización. Separar recepción del archivo de acceso automatizado y permisos documentados.</li>
<li><strong>Adaptador institucional independiente:</strong> conservar campos originales, seudonimizar PIDM
con mecanismo institucional estable y trabajar en un espacio local privado. No reemplazar los adaptadores OULAD/sintético ni la base pública.</li>
<li><strong>Tablas de período:</strong> crear matrícula por persona–período–programa, resultados por
persona–período–sección, dimensión de programas y calendario. Registrar conflictos y no coincidencias.</li>
<li><strong>Marts propuestos:</strong> <code>mart_student_course_term</code> y
<code>mart_student_term</code>. El contrato y la disciplina de validación existentes se reutilizan;
los marts semanales actuales exigen tablas con semana y no deben recibir semanas inventadas.</li>
<li><strong>Variables:</strong> promedio y dispersión de notas numéricas anteriores, número de resultados observados,
proporción de resultados conocidos, resultados adversos según diccionario, intentos repetidos observados,
tendencia entre períodos comparables, programa/campus/jornada vigentes al corte y condición nuevo/continuo
validada. Carga ponderada solo después de aclarar unidades. Sin PIDM, secuencias administrativas ni resultados futuros como predictores.</li>
<li><strong>Segmentación:</strong> comparar al menos K-means, mezcla gaussiana y clustering jerárquico o bisectriz.
Evaluar tamaños, separación, estabilidad temporal e interpretación por direcciones. No convertir segmentos en niveles de riesgo automáticamente.</li>
<li><strong>Predicción:</strong> comparar regresión logística, random forest y gradient boosting con reglas sencillas.
Usar cortes temporales, calibración, AUC/AP, Brier y captura a capacidad 10/20/30%, con incertidumbre por persona.</li>
<li><strong>Producto:</strong> incorporar una sección institucional FAN separada del benchmark público,
con alcance, cobertura y período visibles; mantener navegación hacia las comparaciones externas.</li>
</ol>
<p><strong>Evaluación temporal propuesta:</strong> para regímenes semestrales, estudiar entrenamiento con
2024-1 → 2024-2, validación con 2024-2 → 2025-1 y prueba con 2025-1 → 2025-2.
Es solo un diseño candidato: requiere confirmar cierres, disponibilidad de notas, cobertura y elegibilidad;
dos años proporcionan pocas transiciones. En estudiantes nuevos no existe historia previa y debe declararse
esa limitación. No hacer partición aleatoria por filas ni usar las notas finales del período objetivo.
La misma persona puede reaparecer longitudinalmente solo si se respeta el tiempo; los identificadores no son variables
y se debe evaluar incertidumbre agrupada por persona y comportamiento en estudiantes nuevos.</p>
<p>Para mayor solidez, solicitar 2022–2023 y 2026, tabla curricular y calendario, diccionario de notas/estados,
fechas de publicación/cambios y egresos/títulos. PAES/NEM/ranking requerirían una fuente institucional vinculable;
no existe en este archivo un puente autorizado hacia los MRUN públicos.</p>
<h2>7. Usabilidad propuesta</h2>
<p>Inicio en FAN con las cuatro sedes del alcance aprobado. Filtros dependientes: nivel/modalidad → régimen
→ año/período → campus → carrera/programa → jornada; permitir separar alumnos nuevos/continuos y mostrar
el criterio de estados efectivos. Online, intercambio, verano y TAV no deben mezclarse por defecto.</p>
<p>Las tarjetas deben distinguir personas únicas, planes, resultados de asignaturas y cobertura de notas.
Un filtro sin observaciones debe mostrar ausencia, nunca volver silenciosamente al universo completo.
Mostrar coma decimal, punto de miles, porcentajes con denominador y “sin dato” distinto de cero.
Las comparaciones entre sedes deben usar carreras, niveles y períodos comparables.</p>
<p>Proponer vistas de Panorama, Asignaturas, Perfiles, Seguimiento entre períodos y Calidad.
Una ficha individual pertenecería a un entorno institucional con control de acceso; los CSV de gestión agregados
llevarían filtros y fecha de corte. Suprimir celdas pequeñas y evitar exportar PIDM o predicciones individuales
desde la aplicación pública.</p>
<h2 id="estado">8. Estado completo del proyecto al 20 de octubre de 2026</h2>"""
status_rows = [
    ["Benchmark externo UA", "Implementado en la solución existente, con despliegue configurado.", "Permanece como contexto complementario. No se comprobó disponibilidad de la URL en esta revisión."],
    ["Arquitectura, contrato y marts semanales", "Implementados y documentados para adaptadores existentes.", "Reutilizar validaciones; añadir la granularidad de período para este extracto."],
    ["Piloto público FAN independiente", "Implementado localmente: perfiles de ingreso, continuidad y titulación; última corrida 8 de octubre.", "Se conserva como evidencia histórica complementaria; aún fuera de Git y del despliegue original."],
    ["Nueva fuente institucional", "Recibida y explorada íntegramente; procedencia universitaria confirmada por el usuario.", "Falta confirmar metadatos de extracción y resolver reglas semánticas. No equivale a acceso automatizado."],
    ["Modelos con notas internas", "Todavía no implementados ni evaluados.", "Ahora son candidatos viables para un piloto entre períodos; no se ha medido su desempeño."],
    ["Student 360 institucional", "No operativo con esta fuente.", "Factible una historia local de resultados 2024–2025 tras integrar y validar."],
    ["Canvas / asistencia / notas parciales", "Siguen ausentes.", "No afirmar alertas semanales ni desenganche conductual real."],
    ["Validación con direcciones", "El registro institucional sigue pendiente: alcance, casos y evaluación de uso.", "Ratificar catálogo, tareas de usuario, decisiones y responsables; no simular su cumplimiento."],
    ["Impacto", "Capacidad técnica y evidencia retrospectiva disponibles; no hay impacto de intervenciones medido.", "Evaluar utilidad para decisiones y luego diseño prospectivo; no atribuir reducción de deserción."],
    ["Documentación y publicación", "README mantiene 'sin accesos' y el despliegue original no integra el piloto ni este Excel.", "Actualizar el estado al incorporar formalmente la fuente; distinguir extracto recibido de conectores disponibles."],
]
body += table(["Componente", "Estado verificado", "Implicación"], status_rows)
body += table(["Resultado previo · datos públicos, sin este Excel", "Resultado de la última corrida"], [
    ["Población de ingreso FAN histórica", previous["quality"]["pilot_entries"]],
    ["Segmentación de ingreso", f"{previous['segments']['winner']}, {previous['segments']['k']} grupos; silueta {decimal(previous['segments']['silhouette'])}"],
    ["Continuidad", f"{previous['cases']['continuidad']['winner']}; AUC {decimal(previous['cases']['continuidad']['test']['auc'])}; captura al 20% {decimal(100*previous['cases']['continuidad']['test']['capture_20'], 1)}%; señal retrospectiva."],
    ["Titulación", f"{previous['cases']['titulacion']['winner']}; AUC {decimal(previous['cases']['titulacion']['test']['auc'])}; experimental."],
])
body += """<p>La población anterior se define por primer ingreso observado y MRUN sobre años anteriores;
no es comparable directamente con la matrícula vigente 2024–2025 por PIDM.
Ninguna de sus métricas mejora automáticamente al recibir el archivo. En esta revisión se verificaron
los diagnósticos de datos y la integridad de la fuente; no se reentrenaron modelos ni se reejecutó toda la suite de la aplicación.</p>
<p><strong>Evaluación:</strong> el proyecto tiene una base técnica sólida y una prueba de concepto local,
y la fuente nueva permite pasar a desempeño académico institucional. Su madurez actual corresponde a
un piloto retrospectivo en preparación de integración. Aún no corresponde declararlo un sistema de alerta
operativo, validado por usuarios o con impacto demostrado.</p>
<h2 id="plan">9. Plan de trabajo propuesto</h2>"""
body += table(["Etapa", "Esfuerzo estimado", "Entregable", "Condición de cierre"], [
    ["1. Acuerdo de datos y alcance", "2–3 días", "Diccionario confirmado, catálogo FAN, estados y calendario; dos casos priorizados.", "Contraparte institucional responde las reglas pendientes."],
    ["2. Integración y calidad", "3–4 días", "Adaptador y dos marts de período, informe de excepciones, trazabilidad y pruebas de cardinalidad.", "Ninguna multiplicación silenciosa; exclusiones explícitas; fechas y etiquetas verificadas."],
    ["3. Diagnóstico y segmentación", "3–5 días", "Tablero académico FAN y comparación de al menos tres alternativas de clustering.", "Segmentos interpretables, tamaños/coberturas suficientes y revisión de direcciones."],
    ["4. Dos pruebas predictivas", "5–8 días", "Tres algoritmos por caso, reglas de referencia, test temporal, calibración y fichas.", "Métricas aceptadas previamente; ningún resultado prometido; si no supera reglas se informa."],
    ["5. Uso y cierre", "3–5 días", "Prueba con representantes de las cuatro sedes, correcciones, guía, actualización y presentación.", "Acta de alcance/aceptación, responsable y plataforma institucional definidos."],
])
body += """<p>Estimación de 3–5 semanas de trabajo, dependiente de respuestas institucionales y disponibilidad del equipo;
no sustituye los hitos formales de la postulación. No es necesario simular datos para el núcleo de estos análisis.
Los ejemplos ficticios pueden mantenerse únicamente para demostrar interacciones.</p>
<h2>10. Próxima solicitud concreta a la universidad</h2>
<ol><li>Fecha de corte, emisor, consulta/filtros y confirmación de cobertura de todos los estados y períodos.</li>
<li>Diccionario de RI, RR, RM, RC, CR, CC, P y REPROB; unidad de CREDIT_HOURS y reglas de cálculo de aprobación.</li>
<li>Estados efectivos de matrícula, definición de las secuencias/prioridad y atribución de cursos a planes simultáneos.</li>
<li>Calendario de períodos con fechas reales, cierres y fechas de disponibilidad de notas/estados.</li>
<li>Catálogo FAN con programas/carreras, vigencias, asignaturas/equivalencias y tratamiento de Online/intercambio.</li>
<li>Historia adicional 2022–2023, actualización 2026 y egresos/títulos para distinguir salidas académicas.</li></ol>
<h2>11. Manejo local y publicación</h2>
<p>El Excel contiene identificadores internos aunque no muestre nombres o RUT. No debe incorporarse al paquete
público de Streamlit. Se comprobó una brecha concreta: el .gitignore actual no excluye este XLSX en data/interim;
Git lo ve como archivo nuevo. El paquete de despliegue usa una lista explícita y actualmente no lo incluye.
Antes de cualquier commit hay que proteger la fuente y ampliar las comprobaciones de exportación para PIDM
y sus nombres prefijados. Esta revisión no hizo commits, pushes ni cambios en el despliegue.</p>
<p>Los análisis y cachés se guardaron en <code>piloto_fan/artifacts/auditoria_excel_20261020/</code>,
ya ignorado por Git. Los scripts de auditoría son independientes. No se modificaron el Excel, el código
de la aplicación original, el puntero de la última corrida del piloto ni los modelos existentes.</p>
<h2 id="variables">12. Inventario completo de los 52 campos</h2>
<p>Tipos y faltantes corresponden al Excel leído. Las interpretaciones administrativas pendientes no se
presentan como un diccionario institucional aprobado. Las claves y descripciones repetidas entre hojas
no son 52 predictores independientes.</p>"""
body += table(["Hoja", "Campo", "Ausentes", "%", "Valores distintos", "Interpretación / uso"], [
    [x["hoja"], x["variable"], x["ausentes"], f"{x['porcentaje_ausentes']:.2f}".replace(".", ","),
     x["valores_distintos"], x["interpretacion"]] for x in inventory])
body += """<footer><p>Archivos de respaldo:
<a href="inventario_y_cobertura.xlsx">inventario y cobertura Excel</a> ·
<a href="diccionario_calidad.html">distribuciones de campos</a> ·
<a href="perfil_completo.json">perfil completo JSON</a> ·
<a href="diagnostico_uniones_y_cobertura.json">diagnóstico de uniones</a> ·
<a href="alcance_fan_propuesto.json">alcance propuesto</a>.
Resultados agregados locales; no publicación autorizada.</p></footer>"""

css = """
body{font:16px/1.55 system-ui,-apple-system,Segoe UI,sans-serif;color:#233247;
background:#f5f7fa;margin:0}header{background:#8b1635;color:white;padding:40px 5vw}
header p{max-width:1050px}.eyebrow{font-size:12px;letter-spacing:1.5px}
main{max-width:1200px;margin:auto;padding:30px;background:white}
h1{font-size:32px;line-height:1.2;max-width:1000px}h2{font-size:23px;margin-top:38px}
.callout{background:#edf4fa;border-left:5px solid #23567d;padding:20px}
nav{margin:24px 0}a{color:#8b1635}table{border-collapse:collapse;width:100%;font-size:14px}
td,th{border:1px solid #d8dfe7;padding:10px;text-align:left;vertical-align:top}
th{background:#edf1f5}tbody tr:nth-child(even){background:#f9fafb}.scroll{overflow-x:auto;margin:18px 0}
code{font-size:12px;overflow-wrap:anywhere}li{margin:10px 0}footer{border-top:2px solid #d8dfe7;margin-top:35px;font-size:13px}
@media print{body,main{background:white}main{padding:0}header{background:white;color:#233247;padding:0}
nav{display:none}h2{break-after:avoid}tr{break-inside:avoid}a{color:inherit}.scroll{overflow:visible}td{overflow-wrap:anywhere}}
"""
(OUT / "informe_nueva_fuente_y_estado.html").write_text(
    '<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" '
    'content="width=device-width, initial-scale=1"><title>Diagnóstico institucional FAN · 20 octubre 2026</title>'
    f"<style>{css}</style></head><body><main>{body}</main></body></html>", encoding="utf-8")
with pd.ExcelWriter(OUT / "inventario_y_cobertura.xlsx", engine="openpyxl") as writer:
    pd.DataFrame(inventory).to_excel(writer, sheet_name="52 variables", index=False)
    pd.DataFrame(quality_rows, columns=["Dimension", "Hallazgo", "Tratamiento"]).to_excel(
        writer, sheet_name="Calidad", index=False)
    pd.DataFrame(fan["by_career"]).to_excel(writer, sheet_name="FAN extracto carreras", index=False)
    pd.DataFrame(core["by_campus"]).to_excel(writer, sheet_name="FAN propuesto campus", index=False)
    pd.DataFrame(core["by_period"]).to_excel(writer, sheet_name="FAN propuesto periodos", index=False)
    pd.DataFrame(status_rows, columns=["Componente", "Estado", "Siguiente paso"]).to_excel(
        writer, sheet_name="Estado proyecto", index=False)
    for ws in writer.book.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                75, max(18, max(len(str(c.value or "")) for c in col[:80]) + 2))
        from openpyxl.styles import Font, PatternFill, Alignment
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="8B1635")
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height = 65 if ws.max_column >= 6 else 45
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if ws.title == "52 variables":
            for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
                row[0].number_format = '0.00"%"'
assert hashlib.file_digest(SOURCE.open("rb"), "sha256").hexdigest() == p["sha256"]
print("Informe HTML e inventario Excel generados. Fuente íntegra; 52 campos; sin identificadores individuales.")
